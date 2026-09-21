# SONGJUNSONG (School of Finance and Economics / Jilin Business and Technology College)
import os
import re
import time
import requests
from datetime import datetime
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from PIL import Image, ImageFile
import io

ImageFile.LOAD_TRUNCATED_IMAGES = True


def get_article_base(include_year: bool = False):
    """文章存储根目录。

    优先级：config.yaml 的 storage.base_dir > 默认（user_dir 下的「文章」）。
    默认值同时适用于容器与 Windows 本地运行：容器里项目根就是 /app，
    解析结果与原先硬编码的 /app/文章 完全一致。
    include_year=True 时再拼上 config.yaml 里的 storage_year 子目录。
    """
    try:
        from werss.config import cfg
        custom = (cfg.get('storage.base_dir', '') or '').strip()
    except Exception:
        custom = ''
    if custom:
        base = os.path.abspath(os.path.join(str(custom), '文章'))
    else:
        try:
            from paths import user_dir
            _root = user_dir()
        except Exception:
            _root = os.path.dirname(os.path.abspath(__file__))
        base = os.path.abspath(os.path.join(_root, '文章'))
    if include_year:
        try:
            from werss.config import cfg
            sy = cfg.get('storage_year', '')
            if sy:
                base = os.path.join(base, str(sy))
        except Exception:
            pass
    return base

class WeChatCrawler:
    def __init__(self):
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
        })

    def _inject_wechat_cookie(self):
        try:
            from werss.driver.token import get as get_token
            cookie_str = get_token("cookie")
            if cookie_str and cookie_str.strip():
                cookie_dict = {}
                for pair in cookie_str.split(';'):
                    pair = pair.strip()
                    if '=' in pair:
                        k, v = pair.split('=', 1)
                        cookie_dict[k.strip()] = v.strip()
                if cookie_dict:
                    self.session.cookies.update(cookie_dict)
                else:
                    print("[Crawler] 警告: cookie 字符串为空或解析失败，可能授权已失效")
                    try:
                        from app.core import debug_log
                        debug_log('WARNING', 'rss_auth', 'cookie 字符串为空或解析失败，RSS 授权可能已失效')
                    except Exception:
                        pass
            else:
                print("[Crawler] 警告: 未获取到 cookie，RSS 授权可能已失效")
                try:
                    from app.core import debug_log
                    debug_log('WARNING', 'rss_auth', '未获取到 cookie，RSS 授权可能已失效')
                except Exception:
                    pass
        except Exception as e:
            print(f"[Crawler] cookie 注入异常: {e}")
            try:
                from app.core import debug_log
                debug_log('ERROR', 'rss_auth', f'cookie 注入异常: {e}')
            except Exception:
                pass

    def new_session(self):
        self.session.close()
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
        })

    def crawl_article(self, url, progress_callback=None, max_retries=3, publish_time=0, auto_save=True):
        """爬取微信公众号文章 — SONGJUNSONG (School of Finance and Economics / Jilin Business and Technology College)"""
        last_error = None

        for attempt in range(max_retries):
            try:
                if progress_callback:
                    progress_callback(10, f'正在获取页面... (尝试 {attempt + 1}/{max_retries})')

                self._inject_wechat_cookie()
                response = self.session.get(url, timeout=15, allow_redirects=True)

                # ---- 显式 HTTP 状态码检测 ----
                status = response.status_code

                # 404: 文章已删除或链接失效 → 无需重试
                if status == 404:
                    print(f"文章不存在 (404): {url[:80]}")
                    return {
                        'success': False,
                        'error': '文章已被删除 (404)',
                        'is_404': True
                    }

                # 410 Gone: 资源永久移除
                if status == 410:
                    print(f"文章已永久移除 (410): {url[:80]}")
                    return {
                        'success': False,
                        'error': '文章已永久移除 (410)',
                        'is_404': True
                    }

                # 403 Forbidden: 可能触发反爬或权限限制 → 标记为授权错误，允许上层触发刷新
                if status == 403:
                    err_msg = f'访问被拒绝 (403)'
                    print(f"[Crawler] {err_msg}: {url[:80]}")
                    # 403 可能是 cookie 过期，标记为授权错误供上层检测
                    return {
                        'success': False,
                        'error': err_msg,
                        'is_403': True,
                        'is_auth_error': True
                    }

                # ---- HTTP 重定向检测 ----
                if response.history:
                    redirect_chain = ' → '.join(
                        [f"{r.status_code} {r.url[:60]}" for r in response.history]
                    )
                    final_url = response.url
                    if final_url != url:
                        print(f"重定向检测: {redirect_chain} → 最终 {final_url[:80]}")
                        # 检查是否被重定向到非微信域名
                        if 'mp.weixin.qq.com' not in final_url:
                            return {
                                'success': False,
                                'error': f'页面重定向到外部域名: {final_url[:100]}',
                                'is_redirect': True
                            }

                # 其他非 2xx 状态码
                if status >= 400:
                    print(f"HTTP 错误 ({status}): {url[:80]}")
                    return {
                        'success': False,
                        'error': f'HTTP 状态码 {status}',
                        'is_http_error': True
                    }

                # ---- 继续正常解析 ----

                if progress_callback:
                    progress_callback(20, '正在解析文章内容...')

                html = response.text
                soup = BeautifulSoup(html, 'html.parser')

                # ---- 反爬验证页面检测 ----
                body_text = soup.get_text() or ''
                anti_crawl_keywords = [
                    '当前环境异常，完成验证后即可继续访问',
                    '该内容已被发布者删除',
                    'The content has been deleted by the author',
                    '内容审核中',
                    '该内容暂时无法查看',
                    '违规无法查看',
                    'Unable to view this content because it violates regulation',
                    '发送失败无法查看',
                ]
                for kw in anti_crawl_keywords:
                    if kw in body_text:
                        print(f"[Crawler] 反爬/异常页面检测: {kw}")
                        try:
                            from app.core import debug_log
                            debug_log('WARNING', 'crawl', f'反爬/异常页面: {kw}')
                        except Exception:
                            pass
                        return {
                            'success': False,
                            'error': kw,
                            'is_anti_crawl': True
                        }

                title = self.extract_title(soup)

                if '【转载】' in title:
                    print(f"检测到转载文章，跳过: {title}")
                    return {
                        'success': False,
                        'error': '转载文章，已跳过',
                        'is_repost': True
                    }

                publish_date = self.extract_publish_date(soup, external_ts=publish_time)
                content = self.extract_content(soup, progress_callback)

                print(f"提取到的信息:")
                print(f"标题: {title}")
                print(f"发布日期: {publish_date.strftime('%Y年%m月%d日')}")
                print(f"内容项数量: {len(content)}")

                if progress_callback:
                    progress_callback(80, '正在生成Word文档...')

                if auto_save:
                    file_path = self.create_word_document(title, publish_date, content, url=url)
                else:
                    file_path = ''

                if progress_callback:
                    progress_callback(100, '完成！')

                return {
                    'success': True,
                    'title': title,
                    'publish_date': publish_date,
                    'content_items': content,
                    'file_path': file_path
                }

            except requests.exceptions.RequestException as e:
                last_error = e
                msg = f"请求失败 (尝试 {attempt + 1}/{max_retries}): {str(e)[:100]}"
                print(msg)
                if progress_callback:
                    progress_callback(10, msg)

                if attempt < max_retries - 1:
                    print(f"等待 2 秒后重试...")
                    time.sleep(2)
                else:
                    print(f"已达到最大重试次数，放弃")

            except Exception as e:
                last_error = e
                msg = f"爬取异常 (尝试 {attempt + 1}/{max_retries}): {str(e)[:100]}"
                print(msg)
                if progress_callback:
                    progress_callback(10, msg)

                if attempt < max_retries - 1:
                    print(f"等待 2 秒后重试...")
                    time.sleep(2)
                else:
                    print(f"已达到最大重试次数，放弃")

        return {
            'success': False,
            'error': str(last_error)[:200] if last_error else '未知错误'
        }

    def extract_title(self, soup):
        title_elem = soup.find('h1', class_='rich_media_title')
        if title_elem:
            return title_elem.get_text().strip()
        return "未知标题"

    def extract_publish_date(self, soup, external_ts=0):
        if external_ts and external_ts > 0:
            try:
                dt = datetime.fromtimestamp(external_ts)
                if dt.year >= 2010 and dt <= datetime.now():
                    print(f"使用API提供的发布日期: {dt.strftime('%Y年%m月%d日')}")
                    return dt
            except Exception:
                pass

        print("开始提取文章发布日期...")

        meta_tags = soup.find_all('meta')
        for meta in meta_tags:
            if meta.get('property') in ['article:published_time', 'article:modified_time'] or \
               meta.get('name') in ['publishdate', 'date', 'pubdate']:
                content = meta.get('content', '')
                if content:
                    print(f"在meta标签中找到日期: {content}")
                    parsed_date = self.parse_date_from_text(content)
                    if parsed_date:
                        print(f"成功解析meta标签中的日期: {parsed_date.strftime('%Y年%m月%d日')}")
                        return parsed_date

        script_tags = soup.find_all('script')
        for script in script_tags:
            if script.string:
                script_content = script.string
                date_patterns = [
                    r'publish_time["\']?\s*[:=]\s*["\']?(\d{4}-\d{1,2}-\d{1,2})',
                    r'createTime["\']?\s*[:=]\s*["\']?(\d{4}-\d{1,2}-\d{1,2})',
                    r'pubDate["\']?\s*[:=]\s*["\']?(\d{4}-\d{1,2}-\d{1,2})',
                    r'date["\']?\s*[:=]\s*["\']?(\d{4}-\d{1,2}-\d{1,2})',
                    r'publish_time["\']?\s*[:=]\s*["\']?(\d{10,13})',
                    r'createTime["\']?\s*[:=]\s*["\']?(\d{10,13})',
                    r'(\d{4}年\d{1,2}月\d{1,2}日)',
                    r'(\d{4}/\d{1,2}/\d{1,2})',
                    r'(\d{1,2}/\d{1,2}/\d{4})',
                ]

                for pattern in date_patterns:
                    matches = re.findall(pattern, script_content)
                    for match in matches:
                        print(f"在脚本中找到日期: {match}")
                        parsed_date = self.parse_date_from_text(match)
                        if parsed_date:
                            print(f"成功解析脚本中的日期: {parsed_date.strftime('%Y年%m月%d日')}")
                            return parsed_date

        date_selectors = [
            ('em', {'id': 'publish_time'}),
            ('span', {'class': 'rich_media_meta_text'}),
            ('em', {'class': 'rich_media_meta_text'}),
            ('span', {'id': 'publish_time'}),
            ('div', {'class': 'rich_media_meta'}),
            ('span', {'class': 'rich_media_meta_nickname'}),
            ('time', {}),
            ('span', {'class': re.compile(r'.*time.*|.*date.*', re.I)}),
            ('div', {'class': re.compile(r'.*time.*|.*date.*', re.I)}),
        ]

        for tag, attrs in date_selectors:
            if isinstance(attrs.get('class'), str) and hasattr(re, 'compile'):
                elements = soup.find_all(tag, class_=attrs['class'])
            else:
                elements = soup.find_all(tag, attrs) if attrs else soup.find_all(tag)

            for elem in elements:
                if elem:
                    date_text = elem.get_text().strip()
                    print(f"找到可能的日期文本: {date_text}")
                    parsed_date = self.parse_date_from_text(date_text)
                    if parsed_date:
                        print(f"成功解析日期: {parsed_date.strftime('%Y年%m月%d日')}")
                        return parsed_date

        print("警告: 无法找到文章发布日期，使用当前日期")
        return datetime.now()

    def parse_date_from_text(self, date_text):
        if not date_text:
            return None

        date_text = date_text.strip()

        if date_text.isdigit():
            try:
                timestamp = int(date_text)
                if timestamp > 1000000000000:
                    timestamp = timestamp / 1000
                elif timestamp < 1000000000:
                    return None

                parsed_date = datetime.fromtimestamp(timestamp)
                now = datetime.now()
                if parsed_date <= now and parsed_date.year >= 2010:
                    return parsed_date
            except (ValueError, OSError):
                pass

        date_formats = [
            '%Y-%m-%d',
            '%Y/%m/%d',
            '%Y.%m.%d',
            '%Y-%m-%d %H:%M:%S',
            '%Y/%m/%d %H:%M:%S',
            '%Y年%m月%d日',
            '%Y年%m月%d日 %H:%M',
            '%Y年%m月%d日%H:%M:%S',
            '%B %d, %Y',
            '%b %d, %Y',
            '%d %B %Y',
            '%d %b %Y',
            '%m/%d/%Y',
            '%d/%m/%Y',
            '%m-%d-%Y',
            '%d-%m-%Y',
        ]

        for fmt in date_formats:
            try:
                parsed_date = datetime.strptime(date_text, fmt)
                now = datetime.now()
                if parsed_date <= now and parsed_date.year >= 2010:
                    return parsed_date
            except ValueError:
                continue

        chinese_date_pattern = r'(\d{4})年(\d{1,2})月(\d{1,2})日'
        match = re.search(chinese_date_pattern, date_text)
        if match:
            try:
                year, month, day = match.groups()
                parsed_date = datetime(int(year), int(month), int(day))
                if parsed_date <= datetime.now() and parsed_date.year >= 2010:
                    return parsed_date
            except ValueError:
                pass

        numeric_date_patterns = [
            r'(\d{4})[/-](\d{1,2})[/-](\d{1,2})',
            r'(\d{1,2})[/-](\d{1,2})[/-](\d{4})',
        ]

        for pattern in numeric_date_patterns:
            match = re.search(pattern, date_text)
            if match:
                try:
                    parts = match.groups()
                    if len(parts[0]) == 4:
                        year, month, day = parts
                    else:
                        month, day, year = parts

                    parsed_date = datetime(int(year), int(month), int(day))
                    if parsed_date <= datetime.now() and parsed_date.year >= 2010:
                        return parsed_date
                except ValueError:
                    pass

        return None

    def extract_content(self, soup, progress_callback=None):
        content_elem = soup.find('div', class_='rich_media_content')
        if not content_elem:
            raise Exception("无法找到文章内容")

        content_items = []
        elements = content_elem.find_all(['p', 'img', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6'])

        filter_patterns = [
            r'撰\s*稿[：:].+',
            r'编\s*辑[：:].+',
            r'责任编辑[：:].+',
            r'初\s*审[：:].+',
            r'复\s*审[：:].+',
            r'终\s*审[：:].+'
        ]

        total_elements = len(elements)
        image_count = 0

        # 预收集所有图片 URL，并发下载以加速（替代串行下载）
        img_indices = []  # [(elem_index, img_src), ...]
        for i, elem in enumerate(elements):
            if elem.name == 'img':
                img_src = elem.get('data-src') or elem.get('src')
                if img_src:
                    img_indices.append((i, img_src))

        img_cache = {}  # {img_src: img_data or None}
        if img_indices:
            from concurrent.futures import ThreadPoolExecutor, as_completed
            _MAX_IMG_WORKERS = 6  # 并发下载线程数
            unique_srcs = list({src for _, src in img_indices})
            print(f"[Crawler] 并发下载 {len(unique_srcs)} 张图片（{min(_MAX_IMG_WORKERS, len(unique_srcs))} 线程）")
            with ThreadPoolExecutor(max_workers=min(_MAX_IMG_WORKERS, len(unique_srcs))) as pool:
                future_to_src = {pool.submit(self.download_image, src): src for src in unique_srcs}
                for future in as_completed(future_to_src):
                    src = future_to_src[future]
                    try:
                        img_cache[src] = future.result()
                    except Exception as e:
                        print(f"[Crawler] 图片下载异常 {src}: {e}")
                        img_cache[src] = None

        for i, elem in enumerate(elements):
            if progress_callback and i % 10 == 0:
                progress = 30 + int((i / total_elements) * 40)
                progress_callback(progress, f'正在处理内容元素 {i+1}/{total_elements}...')

            if elem.name == 'img':
                image_count += 1
                img_src = elem.get('data-src') or elem.get('src')
                if img_src:
                    img_data = img_cache.get(img_src)
                    if img_data:
                        content_items.append({
                            'type': 'image',
                            'data': img_data,
                            'index': image_count
                        })
            else:
                text = elem.get_text().strip()
                if text:
                    should_filter = False
                    for pattern in filter_patterns:
                        if re.search(pattern, text):
                            should_filter = True
                            print(f"过滤编辑信息: {text}")
                            break

                    if not should_filter:
                        content_items.append({
                            'type': 'text',
                            'content': text,
                            'tag': elem.name
                        })

        if content_items:
            last_image_index = -1
            for i in range(len(content_items) - 1, -1, -1):
                if content_items[i]['type'] == 'image':
                    last_image_index = i
                    break

            if last_image_index >= 0:
                removed_item = content_items.pop(last_image_index)
                print(f"移除最后一张图片 (索引: {last_image_index})")

        return content_items

    def download_image(self, img_url):
        try:
            if not img_url.startswith('http'):
                img_url = 'https:' + img_url

            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36',
                'Referer': 'https://mp.weixin.qq.com/'
            }

            response = requests.get(img_url, headers=headers, timeout=10)
            if response.status_code == 200:
                img_data = response.content
                img_size = len(img_data)

                if img_size < 10 * 1024:
                    print(f"跳过小图片 ({img_size} bytes < 10KB): {img_url}")
                    return None

                print(f"下载图片成功 ({img_size} bytes): {img_url}")
                return img_data
        except Exception as e:
            print(f"下载图片失败: {e}")
        return None

    def create_word_document(self, title, publish_date, content_items, save_file=True, url=''):
        from docx.shared import RGBColor
        import tempfile

        doc = Document()
        self.setup_document_style(doc)
        self.create_custom_styles(doc)

        title_paragraph = doc.add_paragraph()
        title_paragraph.style = 'Title Style'
        title_run = title_paragraph.add_run(title)
        title_run.font.size = Pt(22)
        title_run.font.name = 'SimHei'
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0, 0, 0)
        title_run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
        title_run._element.rPr.rFonts.set(qn('w:ascii'), 'SimHei')
        title_run._element.rPr.rFonts.set(qn('w:hAnsi'), 'SimHei')
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        date_str = publish_date.strftime('%Y年%m月%d日')
        date_paragraph = doc.add_paragraph()
        date_paragraph.style = 'Date Style'
        date_run = date_paragraph.add_run(date_str)
        date_run.font.size = Pt(15)
        date_run.font.name = '仿宋_GB2312'
        date_run.font.color.rgb = RGBColor(0, 0, 0)
        date_run._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋_GB2312')
        date_run._element.rPr.rFonts.set(qn('w:ascii'), '仿宋_GB2312')
        date_run._element.rPr.rFonts.set(qn('w:hAnsi'), '仿宋_GB2312')
        date_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_paragraph()

        temp_files = []

        for item in content_items:
            if item['type'] == 'text':
                p = doc.add_paragraph()
                p.style = 'Body Style'
                run = p.add_run(item['content'])
                run.font.size = Pt(15)
                run.font.name = '仿宋_GB2312'
                run.font.color.rgb = RGBColor(0, 0, 0)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋_GB2312')
                run._element.rPr.rFonts.set(qn('w:ascii'), '仿宋_GB2312')
                run._element.rPr.rFonts.set(qn('w:hAnsi'), '仿宋_GB2312')
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            elif item['type'] == 'image':
                try:
                    img_data = item['data']

                    with Image.open(io.BytesIO(img_data)) as img:
                        img_w, img_h = img.size
                        img_fmt = img.format or 'PNG'
                    if img_w > img_h:
                        width_inches = Inches(6)
                    else:
                        width_inches = Inches(6 * img_w / img_h)

                    ext = '.' + img_fmt.lower()
                    fd, tmp_path = tempfile.mkstemp(suffix=ext)
                    with os.fdopen(fd, 'wb') as tmpf:
                        tmpf.write(img_data)
                    temp_files.append(tmp_path)

                    doc.add_picture(tmp_path, width=width_inches)

                except Exception as e:
                    print(f"插入图片失败: {e}")

        if save_file:
            result = self.save_document(doc, title, publish_date, url)
            for tmp in temp_files:
                try:
                    os.remove(tmp)
                except Exception:
                    pass
        else:
            doc._temp_image_files = temp_files
            result = doc

        return result

    def create_custom_styles(self, doc):
        styles = doc.styles

        if 'Title Style' not in [s.name for s in styles]:
            title_style = styles.add_style('Title Style', WD_STYLE_TYPE.PARAGRAPH)
            title_style.font.size = Pt(22)
            title_style.font.name = 'SimHei'
            title_style.font.bold = True
            title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if 'Date Style' not in [s.name for s in styles]:
            date_style = styles.add_style('Date Style', WD_STYLE_TYPE.PARAGRAPH)
            date_style.font.size = Pt(15)
            date_style.font.name = '仿宋_GB2312'
            date_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        if 'Body Style' not in [s.name for s in styles]:
            body_style = styles.add_style('Body Style', WD_STYLE_TYPE.PARAGRAPH)
            body_style.font.size = Pt(15)
            body_style.font.name = '仿宋_GB2312'
            body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def generate_preview_html(self, title, publish_date, content_items):
        import base64
        html_content = f"""
        <div class="document-preview">
            <div class="title" style="font-family: SimHei, '黑体', sans-serif; font-size: 22pt; font-weight: bold; text-align: center; margin-bottom: 20px;">
                {title}
            </div>
            <div class="date" style="font-family: '仿宋_GB2312', '仿宋', serif; font-size: 15pt; text-align: right; margin-bottom: 30px;">
                {publish_date.strftime('%Y年%m月%d日')}
            </div>
            <div class="content">
        """

        for item in content_items:
            if item['type'] == 'text':
                html_content += f"""
                <p style="font-family: '仿宋_GB2312', '仿宋', serif; font-size: 15pt; text-align: justify; line-height: 1.5; margin-bottom: 15px;">
                    {item['content']}
                </p>
                """
            elif item['type'] == 'image':
                try:
                    img_base64 = base64.b64encode(item['data']).decode('utf-8')
                    html_content += f"""
                    <div style="text-align: center; margin: 20px 0;">
                        <img src="data:image/jpeg;base64,{img_base64}" style="max-width: 100%; height: auto;" />
                    </div>
                    """
                except Exception as e:
                    print(f"预览图片失败: {e}")

        html_content += """
            </div>
        </div>
        """

        return html_content

    def setup_document_style(self, doc):
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

    def _get_article_base(self):
        """读取存储年份配置，拼装文章存储根路径"""
        return get_article_base(include_year=True)

    def save_document(self, doc, title, publish_date, url=''):
        base_path = self._get_article_base()

        print(f"使用发布日期: {publish_date}")
        date_folder = publish_date.strftime('%Y%m%d')
        folder_path = os.path.join(base_path, date_folder)

        os.makedirs(folder_path, exist_ok=True)
        print(f"创建文件夹: {folder_path} (基于发布日期: {publish_date.strftime('%Y年%m月%d日')})")

        clean_title = re.sub(r'[<>:"/\\|?*]', '', title)
        clean_title = clean_title.replace(' ', '_')
        clean_title = clean_title.replace('\n', '_')
        clean_title = clean_title.replace('\r', '_')
        if len(clean_title) > 100:
            clean_title = clean_title[:100]

        filename = f"{date_folder}_{clean_title}.docx"
        file_path = os.path.join(folder_path, filename)

        print(f"保存文档到: {file_path}")
        print(f"文件名包含发布日期: {date_folder} ({publish_date.strftime('%Y年%m月%d日')})")

        doc.save(file_path)

        if os.path.exists(file_path):
            file_size = os.path.getsize(file_path)
            print(f"文档保存成功，大小: {file_size} bytes")
            print(f"最终保存路径: {file_path}")
            return file_path
        else:
            raise Exception(f"文档保存失败: {file_path}")

    def __del__(self):
        if hasattr(self, 'session') and self.session:
            self.session.close()
